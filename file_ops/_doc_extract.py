"""Real document-text extraction for read_file's document-extension branch.

Used by both the old and new read paths via mocks.extract_document_bytes.
Kept as its own module (not folded into mocks.py) so openpyxl/anydoc are
only imported when a matching extension is actually read, not on every
mocks.py import.

Two kinds of extractor (Plan07, owner 2026-10-05):
- OURS: .xlsx (sheet names kept - anydoc drops them), .ipynb (cells
  labelled, outputs shown).
- anydoc (`firecrawl-anydoc`, a plain PyPI dependency of ours): PDF, Office
  incl. .docx/.pptx, OpenDocument, RTF, EPUB. .docx moved there after a
  comparison (headings, lists, nested tables, escaped `|` - ours lost them).
  Called directly, NOT through Hermes' tools/read_extract.py - no foreign
  seam. The extension list is a snapshot of Hermes' ANYDOC_EXTENSIONS on 54bc5e509c; anydoc sniffs the
  real format from the bytes, as Hermes' own byte path does.
No OCR: scanned PDF pages come back as an ERROR with a recovery hint.
SQLite is deliberately not a "document" - agents read databases with Python.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

OWN_EXTENSIONS = frozenset({".ipynb", ".xlsx"})
ANYDOC_EXTENSIONS = frozenset({
    ".pdf", ".doc", ".docx", ".docm", ".rtf", ".epub",
    ".ppt", ".pps", ".pot", ".pptx", ".pptm", ".ppsx", ".ppsm",
    ".xls", ".xlsm", ".xlsb",
    ".odt", ".ods", ".odp",
})
EXTRACTABLE_EXTENSIONS = OWN_EXTENSIONS | ANYDOC_EXTENSIONS

ANYDOC_REQUIREMENT = "firecrawl-anydoc>=0.2.4"


class ExtractionError(Exception):
    """Extraction failed; the message is shown to the model as-is."""


def extract(data: bytes, path: str):
    """Return plain text for a supported document extension, or None."""
    ext = Path(path).suffix.lower()
    if ext == ".ipynb":
        return _extract_ipynb(data)
    if ext == ".xlsx":
        return _extract_xlsx(data)
    if ext in ANYDOC_EXTENSIONS:
        return _extract_anydoc(data, path)
    return None


# ── anydoc ──────────────────────────────────────────────────────────

def _extract_anydoc(data: bytes, path: str) -> str:
    ext = Path(path).suffix.lower()
    try:
        import anydoc
    except ImportError:
        raise ExtractionError(
            f"{ext} needs the anydoc converter, which is not installed "
            f"(pip install '{ANYDOC_REQUIREMENT}'), or convert the file yourself "
            "via terminal (e.g. libreoffice --headless --convert-to txt).") from None
    try:
        text = anydoc.to_markdown_bytes(data)
    except Exception as exc:
        needs_ocr = getattr(anydoc, "NeedsOcrError", None)
        if needs_ocr is not None and isinstance(exc, needs_ocr):
            raise ExtractionError(_scanned_pdf_message(path, getattr(exc, "pages", None))) from None
        raise ExtractionError(f"{type(exc).__name__}: {exc}") from None
    if not isinstance(text, str) or not text.strip():
        raise ExtractionError("document contains no extractable text")
    return text.rstrip("\n")


def _scanned_pdf_message(path: str, pages) -> str:
    page_list = ", ".join(str(p) for p in pages) if pages else "unknown"
    return (
        f"scanned pages without a text layer: {page_list} - there is no text to "
        "extract. To see them: render those pages to images "
        f"(pdftoppm -jpeg -r 150 -f <first> -l <last> \"{path}\" <out_prefix>) "
        "and look at them with vision_analyze.")


# ── .ipynb ──────────────────────────────────────────────────────────

_CELL_LABELS = {"markdown": "Markdown", "code": "Code", "raw": "Raw"}
_MAX_OUTPUT_CHARS = 20_000  # per code cell: one runaway log must not flood the read


def _text(value) -> str:
    if isinstance(value, list):
        return "".join(v for v in value if isinstance(v, str))
    return value if isinstance(value, str) else ""


def _output_text(output) -> str:
    """One cell output as text: streams, errors and plain-text results kept;
    images/HTML/widgets become a one-line placeholder."""
    if not isinstance(output, dict):
        return ""
    otype = output.get("output_type")
    if otype == "stream":
        return _text(output.get("text"))
    if otype in ("error", "pyerr"):
        return f"Error: {output.get('ename', '')}: {output.get('evalue', '')}"
    if otype not in ("execute_result", "display_data", "pyout"):
        return ""
    data = output.get("data")
    if not isinstance(data, dict):  # nbformat v3: text sits flat on the output
        data = {"text/plain": output.get("text")}
    for mime in ("text/plain", "text/markdown"):
        body = _text(data.get(mime))
        if body.strip():
            return body
    kinds = ", ".join(str(m) for m in data) or "unknown"
    return f"[{kinds} output - omitted]"


def _extract_ipynb(data: bytes):
    try:
        nb = json.loads(data.decode("utf-8-sig", errors="replace"))
    except json.JSONDecodeError:
        return None
    if not isinstance(nb, dict):
        return None
    cells = nb.get("cells")
    if not isinstance(cells, list):  # nbformat v3: cells live under worksheets
        cells = [c for ws in nb.get("worksheets", []) if isinstance(ws, dict)
                 for c in ws.get("cells", [])]
    counts = dict.fromkeys(_CELL_LABELS, 0)
    out = []
    for cell in cells:
        typ = cell.get("cell_type") if isinstance(cell, dict) else None
        if typ not in _CELL_LABELS:
            continue
        counts[typ] += 1
        source = _text(cell.get("source", cell.get("input", ""))).rstrip("\n")
        out += [f"# -- {_CELL_LABELS[typ]} cell {counts[typ]} --", source, ""]
        if typ == "code":
            outputs = cell.get("outputs")
            joined = "\n".join(filter(None, map(_output_text, outputs))) if isinstance(outputs, list) else ""
            joined = joined.rstrip("\n")
            if len(joined) > _MAX_OUTPUT_CHARS:
                cut = len(joined) - _MAX_OUTPUT_CHARS
                joined = joined[:_MAX_OUTPUT_CHARS] + f"\n... [{cut:,} output chars truncated]"
            if joined.strip():
                out += [f"# -- Output (cell {counts[typ]}) --", joined, ""]
    return "\n".join(out).strip()


# ── .xlsx ──────────────────────────────────────────────────────────

def _render_grid(rows) -> str:
    """Plain pipe-separated grid - not a markdown table (no alignment dashes):
    nothing renders this output as markdown, and the dashes would just be
    noise a model has to parse back out."""
    return "\n".join(" | ".join("" if c is None else str(c) for c in row) for row in rows)


def _extract_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = [row for row in ws.iter_rows(values_only=True)
                if row and any(c is not None for c in row)]
        if rows:
            parts.append(f"[Sheet: {ws.title}]\n" + _render_grid(rows))
        else:
            parts.append(f"[Sheet: {ws.title}] (empty)")
    return "\n\n".join(parts)
